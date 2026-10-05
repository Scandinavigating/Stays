import os
import re
import pandas as pd
import openpyxl

# ================= PATH CONFIGURATIE =================
# CONVERSION_DIR = '.../Scandinavigating/Conversion Page'
CONVERSION_DIR = os.path.dirname(os.path.abspath(__file__))
SCANDINAVIGATING_ROOT = os.path.dirname(CONVERSION_DIR)

EXCEL_PATH = os.path.join(SCANDINAVIGATING_ROOT, "Uitvoerend", "Generated Videos.xlsx")
THUMBNAILS_DIR = os.path.join(CONVERSION_DIR, "Thumbnails")
OUTPUT_HTML = os.path.join(CONVERSION_DIR, "index.html")

TARGET_SHEET = "Generated Videos"
TARGET_TABLE = "FinalContentPlan"
GENERAL_MAP_URL = "https://www.stay22.com/embed/gm?aid=scandinavigating&lat=64.0&lng=16.0&zoom=4"

def load_excel_data():
    """
    Leest de Excel uit en haalt specifiek de tabel 'FinalContentPlan' op 
    uit het werkblad 'Generated Videos'.
    """
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"Excel-bestand niet gevonden op: {EXCEL_PATH}")

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    if TARGET_SHEET not in wb.sheetnames:
        raise ValueError(f"Werkblad '{TARGET_SHEET}' niet gevonden in {EXCEL_PATH}")
        
    ws = wb[TARGET_SHEET]

    # Zoek naar de officiële Excel-tabel 'FinalContentPlan'
    table_range = None
    for tbl in ws.tables.values():
        if tbl.name.strip().lower() == TARGET_TABLE.lower():
            table_range = tbl.ref
            break

    # Als de tabelnaam niet direct gevonden wordt, lees het hele werkblad
    if not table_range:
        print(f"[!] Tabel '{TARGET_TABLE}' niet expliciet gedefinieerd; fallback naar hele sheet.")
        df = pd.read_excel(EXCEL_PATH, sheet_name=TARGET_SHEET)
    else:
        # Lees het exacte bereik van de tabel
        min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(table_range)
        data = []
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True):
            data.append(list(row))
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(data[0])]
        df = pd.DataFrame(data[1:], columns=headers)

    # Zorg dat kolom 'ID' als string te vergelijken is
    if "ID" in df.columns:
        df["ID_clean"] = df["ID"].astype(str).str.strip()
    else:
        df["ID_clean"] = ""

    return df

def clean_thumbnail_title(filename):
    """
    Haalt nummers zoals '1. ' aan het begin weg en haalt '(video_xx)' weg.
    Bijv: '1. Rovaniemi, Finland (video_03).jpeg' -> 'Rovaniemi, Finland'
    """
    name_no_ext = os.path.splitext(filename)[0]
    # Verwijder volgnummer aan het begin (bijv. "1. " of "01 - ")
    name_no_num = re.sub(r"^\d+[\.\s\-_]+", "", name_no_ext)
    # Verwijder alles tussen haakjes inclusief de haakjes (bijv. "(video_03)")
    cleaned_name = re.sub(r"\(.*?\)", "", name_no_num)
    return cleaned_name.strip(" ,.-_")

def extract_id_from_filename(filename):
    """
    Haalt de tekst tussen ronde haakjes op.
    Bijv: '1. Rovaniemi, Finland (video_03).jpeg' -> 'video_03'
    """
    match = re.search(r"\((.*?)\)", filename)
    if match:
        return match.group(1).strip()
    return None

def build_grid_cards(df):
    """
    Scant de Thumbnails map, koppelt via de ID aan de Excel data,
    en selecteert de Roam URL (primair) of Map URL (secundair).
    """
    if not os.path.exists(THUMBNAILS_DIR):
        os.makedirs(THUMBNAILS_DIR, exist_ok=True)
        return []

    # Filter op geldige afbeeldingsbestanden
    valid_exts = {".jpg", ".jpeg", ".png", ".webp"}
    files = [f for f in os.listdir(THUMBNAILS_DIR) if os.path.splitext(f)[1].lower() in valid_exts]

    # Sorteer op eventueel volgnummer vooraan
    def sort_key(name):
        m = re.match(r"^(\d+)", name)
        return int(m.group(1)) if m else 9999
    files.sort(key=sort_key)

    cards = []

    for fname in files:
        video_id = extract_id_from_filename(fname)
        display_title = clean_thumbnail_title(fname)
        image_src = f"Thumbnails/{fname}"
        
        target_link = ""

        if video_id:
            # Zoek match in Excel
            match_row = df[df["ID_clean"].str.lower() == video_id.lower()]
            if not match_row.empty:
                row = match_row.iloc[0]
                
                # Prioriteit 1: Stay22_Roam_URL (Allez Roam)
                roam_url = str(row.get("Stay22_Roam_URL", "")).strip()
                # Prioriteit 2: Stay22_Map_URL (Embed GM)
                map_url = str(row.get("Stay22_Map_URL", "")).strip()

                if roam_url and roam_url.lower() != "nan" and roam_url.startswith("http"):
                    target_link = roam_url
                elif map_url and map_url.lower() != "nan" and map_url.startswith("http"):
                    target_link = map_url

        # Als er geen link gevonden kon worden in Excel, gebruik de algemene kaart als fallback
        if not target_link:
            target_link = GENERAL_MAP_URL
            print(f"[!] Geen geldige URL gevonden in Excel voor '{fname}' (ID: {video_id}).")

        cards.append({
            "title": display_title,
            "link": target_link,
            "image": image_src
        })

    return cards

def generate_conversion_page():
    try:
        df = load_excel_data()
    except Exception as e:
        print(f"Fout bij openen van Excel: {e}")
        return

    cards = build_grid_cards(df)

    html_code = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Scandinavigating | Travel Guides, Stays & Maps</title>
  <style>
    * {{
      box-sizing: border-box;
      margin: 0;
      padding: 0;
    }}
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      background-color: #0b1325;
      color: #ffffff;
      padding: 24px 12px 60px 12px;
      display: flex;
      flex-direction: column;
      align-items: center;
    }}
    .header {{
      text-align: center;
      max-width: 500px;
      margin-bottom: 24px;
    }}
    .header h1 {{
      font-size: 1.35rem;
      letter-spacing: 1.5px;
      text-transform: uppercase;
      font-weight: 800;
      margin-bottom: 6px;
    }}
    .header p {{
      color: #94a3b8;
      font-size: 0.88rem;
      line-height: 1.4;
    }}
    .master-btn {{
      display: inline-block;
      margin-top: 16px;
      padding: 12px 24px;
      background-color: #56cfe1;
      color: #0b1325;
      font-weight: 700;
      font-size: 0.88rem;
      border-radius: 30px;
      text-decoration: none;
      box-shadow: 0 4px 15px rgba(86, 207, 225, 0.28);
      transition: transform 0.15s ease, background-color 0.15s ease;
    }}
    .master-btn:hover {{
      background-color: #48b8c9;
    }}
    .master-btn:active {{
      transform: scale(0.97);
    }}
    .grid-container {{
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 6px;
      width: 100%;
      max-width: 600px;
      margin-top: 12px;
    }}
    .grid-item {{
      position: relative;
      aspect-ratio: 9 / 16;
      border-radius: 6px;
      overflow: hidden;
      background-color: #1e293b;
      text-decoration: none;
      box-shadow: 0 2px 8px rgba(0, 0, 0, 0.4);
    }}
    .grid-item img {{
      width: 100%;
      height: 100%;
      object-fit: cover;
      display: block;
      transition: transform 0.25s ease;
    }}
    .grid-item:hover img, .grid-item:active img {{
      transform: scale(1.04);
    }}
    .overlay {{
      position: absolute;
      bottom: 0;
      left: 0;
      right: 0;
      background: linear-gradient(to top, rgba(11, 19, 37, 0.95) 0%, rgba(11, 19, 37, 0.4) 65%, transparent 100%);
      padding: 24px 6px 8px 6px;
      font-size: 0.72rem;
      font-weight: 600;
      text-align: center;
      color: #ffffff;
      line-height: 1.25;
      text-shadow: 0 1px 3px rgba(0, 0, 0, 0.8);
    }}
  </style>
</head>
<body>

  <header class="header">
    <h1>SCANDINAVIGATING</h1>
    <p>Tap any video you watched to discover accommodations in the destination you wish to visit, or click the button below to explore the map!</p>
    <a href="{GENERAL_MAP_URL}" class="master-btn" target="_blank">
      🗺️ Explore All Stays
    </a>
  </header>

  <main class="grid-container">
"""

    for card in cards:
        html_code += f"""    <a href="{card['link']}" class="grid-item" target="_blank">
      <img src="{card['image']}" alt="{card['title']}">
      <div class="overlay">📍 {card['title']}</div>
    </a>
"""

    html_code += """  </main>

</body>
</html>
"""

    with open(OUTPUT_HTML, "w", encoding="utf-8") as f:
        f.write(html_code)

    print(f"Succesvol gegenereerd! {len(cards)} tegels toegevoegd aan index.html.")

if __name__ == "__main__":
    generate_conversion_page()